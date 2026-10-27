"""Semiconductor LSI yield analysis and automated PowerPoint report generator."""

import datetime
import glob
import io
import logging
from pathlib import Path
import random
import time
import urllib.parse

# Third-party library imports (configure matplotlib backend prior to pyplot import)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
import seaborn as sns

# Set Japanese font configuration for Windows
try:
    plt.rcParams["font.family"] = "Meiryo"
except Exception:
    pass

# ---------------------------------------------------------
# 0. Initial Path and Directory Setup
# ---------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "inputs"
OUTPUT_DIR = BASE_DIR / "outputs"
LOG_DIR = BASE_DIR / "logs"

for d in [INPUT_DIR, OUTPUT_DIR, LOG_DIR]:
    d.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    filename=LOG_DIR / "app.log",
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)


# ---------------------------------------------------------
# 1. Excel File Loading and Multi-Sheet Validation
# ---------------------------------------------------------
def load_and_validate_yield_excel(file_path: Path) -> pd.DataFrame:
    """Parse and validate yield summary data across all lot sheets."""
    try:
        xls = pd.ExcelFile(file_path)
    except Exception as e:
        logger.error(f"Failed to open Excel file: {e}")
        raise ValueError(f"Failed to load Excel file: {e}") from e

    all_rows = []
    target_sheets = xls.sheet_names

    # 1. Overall aggregation from 'Summary_Overview' sheet
    if "Summary_Overview" in target_sheets:
        try:
            df_sum = pd.read_excel(
                file_path, sheet_name="Summary_Overview", header=None
            )
            header_idx = 2
            for r in range(min(5, len(df_sum))):
                row_vals = df_sum.iloc[r].astype(str).values
                if any("lot_no" in str(v).lower() for v in row_vals):
                    header_idx = r
                    break

            df_sum = pd.read_excel(
                file_path, sheet_name="Summary_Overview", header=header_idx
            )
            df_sum.columns = [
                str(c).strip().replace("\n", "_") for c in df_sum.columns
            ]

            for _, row in df_sum.iterrows():
                lot_val = None
                for col in df_sum.columns:
                    if "lot" in col.lower():
                        lot_val = str(row[col]).strip()
                        break
                if not lot_val or lot_val.lower() in [
                    "nan",
                    "none",
                    "",
                    "nan.0",
                ]:
                    continue

                pass_cnt, fail_cnt, target_val = np.nan, np.nan, np.nan
                for col in df_sum.columns:
                    c_low = col.lower()
                    if "pass" in c_low:
                        pass_cnt = pd.to_numeric(row[col], errors="coerce")
                    elif "fail" in c_low:
                        fail_cnt = pd.to_numeric(row[col], errors="coerce")
                    elif "target" in c_low:
                        target_val = pd.to_numeric(row[col], errors="coerce")

                if pd.isna(pass_cnt) or pd.isna(fail_cnt):
                    continue

                total_pcs = int(pass_cnt + fail_cnt)
                yield_rate = (
                    pass_cnt / total_pcs * 100 if total_pcs > 0 else 0
                )

                all_rows.append(
                    {
                        "LOT_No": lot_val,
                        "pass_count": int(pass_cnt),
                        "fail_count": int(fail_cnt),
                        "total_pcs": total_pcs,
                        "yield_rate": yield_rate,
                        "target_yield": (
                            target_val if not pd.isna(target_val) else np.nan
                        ),
                    }
                )
        except Exception as e:
            logger.warning(f"Failed to load Summary_Overview sheet: {e}")

    # 2. Fallback parsing from individual 'LOT_NO_*' sheets
    if not all_rows:
        for sname in target_sheets:
            if sname.startswith("LOT_NO_"):
                try:
                    df_s = pd.read_excel(
                        file_path, sheet_name=sname, header=None
                    )
                    lot_no = (
                        str(df_s.iloc[2, 1]).strip()
                        if len(df_s) > 2
                        else sname.replace("LOT_NO_", "")
                    )
                    pass_cnt = pd.to_numeric(df_s.iloc[2, 4], errors="coerce")
                    fail_cnt = pd.to_numeric(df_s.iloc[2, 5], errors="coerce")
                    target_val = pd.to_numeric(
                        df_s.iloc[2, 6], errors="coerce"
                    )

                    if pd.notna(pass_cnt) and pd.notna(fail_cnt):
                        total_pcs = int(pass_cnt + fail_cnt)
                        yield_rate = (
                            pass_cnt / total_pcs * 100 if total_pcs > 0 else 0
                        )
                        all_rows.append(
                            {
                                "LOT_No": lot_no,
                                "pass_count": int(pass_cnt),
                                "fail_count": int(fail_cnt),
                                "total_pcs": total_pcs,
                                "yield_rate": yield_rate,
                                "target_yield": (
                                    target_val
                                    if not pd.isna(target_val)
                                    else np.nan
                                ),
                            }
                        )
                except Exception as ex:
                    logger.warning(f"Skipping sheet {sname} due to error: {ex}")

    return pd.DataFrame(all_rows)


# ---------------------------------------------------------
# 2. Visualization Module
# ---------------------------------------------------------
def generate_trend_chart(df: pd.DataFrame, output_path: Path):
    """Generate overall yield trend line chart."""
    plt.figure(figsize=(8, 4.2), dpi=300)

    sns.lineplot(
        data=df,
        x="LOT_No",
        y="yield_rate",
        marker="o",
        color="#0284c7",
        linewidth=2.5,
        label="実績歩留まり",
        sort=False,
    )

    if "target_yield" in df.columns and not df["target_yield"].isna().all():
        t_vals = df["target_yield"].apply(
            lambda x: x * 100 if x <= 1.0 else x
        )
        plt.plot(
            df["LOT_No"],
            t_vals,
            color="#2563eb",
            marker="x",
            linewidth=2.0,
            linestyle="--",
            label="Target歩留まり",
        )

    plt.title("歩留まりトレンド", fontsize=14, fontweight="bold", pad=15)
    plt.xlabel("LOT No", fontsize=11)
    plt.ylabel("歩留まり (%)", fontsize=11)
    plt.ylim(0, 105)
    plt.grid(True, linestyle="--", alpha=0.5)
    plt.xticks(rotation=30)
    plt.legend(loc="lower right")
    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()


def generate_lot_pareto_charts(
    file_path: Path, output_dir: Path, file_stem: str, timestamp: str
) -> list:
    """Generate wide Pareto analysis charts for each individual lot."""
    pareto_results = []
    xls = pd.ExcelFile(file_path)

    for sname in xls.sheet_names:
        if sname.startswith("LOT_NO_"):
            lot_no = sname.replace("LOT_NO_", "")
            try:
                df_s = pd.read_excel(file_path, sheet_name=sname, header=None)

                tno_list = []
                start_row = -1
                col_tno_idx = -1
                col_val_idx = -1

                # Automatic header index detection for defect breakdown
                for r in range(len(df_s)):
                    row_vals = [str(x).strip() for x in df_s.iloc[r].values]
                    if "不良項目 (Tno)" in row_vals:
                        start_row = r + 1
                        col_tno_idx = row_vals.index("不良項目 (Tno)")
                        for i, v in enumerate(row_vals):
                            if "発生数量" in v:
                                col_val_idx = i
                                break
                        if col_val_idx == -1:
                            col_val_idx = col_tno_idx + 1
                        break

                if start_row != -1:
                    for r in range(start_row, len(df_s)):
                        tno_name = str(
                            df_s.iloc[r, col_tno_idx]
                        ).strip()
                        if tno_name.lower() in [
                            "nan",
                            "none",
                            "",
                        ] or pd.isna(df_s.iloc[r, col_tno_idx]):
                            continue

                        val = pd.to_numeric(
                            df_s.iloc[r, col_val_idx], errors="coerce"
                        )
                        if pd.notna(val) and val > 0:
                            tno_list.append((tno_name, val))

                if not tno_list:
                    tno_list = [("N/A", 1)]

                defect_vals = pd.Series(
                    [val for name, val in tno_list],
                    index=[name for name, val in tno_list],
                )
                defect_vals = defect_vals.sort_values(ascending=False)

                total_defects = (
                    defect_vals.sum() if defect_vals.sum() > 0 else 1
                )
                cum_percentage = (defect_vals.cumsum() / total_defects) * 100

                # Optimized layout size for widescreen PowerPoint (12 x 3.8 inches)
                fig, ax1 = plt.subplots(figsize=(12, 3.8), dpi=300)
                x_indexes = np.arange(len(defect_vals))

                # Bar chart for defect counts
                ax1.bar(
                    x_indexes,
                    defect_vals.values,
                    color="#dc2626",
                    alpha=0.85,
                    width=0.6,
                    label="不良数",
                )
                ax1.set_xlabel(
                    "Fail_Tno (個別テスト番号)", fontsize=11, fontweight="bold"
                )
                ax1.set_ylabel(
                    "不良数 (pcs)", color="#dc2626", fontsize=11, fontweight="bold"
                )
                ax1.tick_params(axis="y", labelcolor="#dc2626")
                ax1.set_xticks(x_indexes)
                ax1.set_xticklabels(
                    defect_vals.index, rotation=35, ha="right", fontsize=9
                )
                ax1.set_ylim(
                    0,
                    defect_vals.max() * 1.15
                    if defect_vals.max() > 0
                    else 10,
                )

                # Line chart for cumulative percentage
                ax2 = ax1.twinx()
                ax2.plot(
                    x_indexes,
                    cum_percentage.values,
                    color="#1e3a8a",
                    marker="s",
                    linewidth=2,
                    label="累積構成比 (%)",
                )
                ax2.axhline(
                    y=80,
                    color="#10b981",
                    linestyle="--",
                    linewidth=1.5,
                    label="80% 基準線",
                )
                ax2.set_ylabel(
                    "累積構成比 (%)", color="#1e3a8a", fontsize=11, fontweight="bold"
                )
                ax2.tick_params(axis="y", labelcolor="#1e3a8a")
                ax2.set_ylim(0, 105)

                plt.title(
                    f"不良要因 パレート分析 LOT_NO.{lot_no}",
                    fontsize=14,
                    fontweight="bold",
                    pad=15,
                )
                plt.tight_layout()

                p_path = (
                    output_dir / f"pareto_{file_stem}_{lot_no}_{timestamp}.png"
                )
                plt.savefig(p_path)
                plt.close()

                # Extract top killer defect information
                top_tno = defect_vals.index[0] if len(defect_vals) > 0 else "N/A"
                top_count = int(defect_vals.iloc[0]) if len(defect_vals) > 0 else 0

                pareto_results.append(
                    {
                        "lot_no": lot_no,
                        "image_path": p_path,
                        "top_tno": top_tno,
                        "top_count": top_count,
                        "total_defects": int(total_defects),
                    }
                )

            except Exception as ex:
                logger.warning(
                    f"Failed to generate Pareto chart for sheet {sname}: {ex}"
                )

    return pareto_results


# ---------------------------------------------------------
# 3. Automated Engineering Analysis Module
# ---------------------------------------------------------
def get_llm_analysis_summary(df: pd.DataFrame) -> dict:
    """Generate overall summary statistics and yield gap description."""
    avg_yield = df["yield_rate"].mean() if not df.empty else 0.0
    worst_lot = df.loc[df["yield_rate"].idxmin()] if not df.empty else None

    return {
        "overall_status": f"全 {len(df)} LOTの平均歩留まりは {avg_yield:.2f}% です。",
        "worst_lot_no": worst_lot["LOT_No"] if worst_lot is not None else "N/A",
        "worst_yield": worst_lot["yield_rate"] if worst_lot is not None else 0.0,
        "trend_comparison": (
            "Target歩留まりに対する歩留まりギャップ（Yield Gap）と、"
            "個別Fail_Tnoごとの不良傾向を確認しました。"
        ),
    }


def get_lot_specific_analysis(row: pd.Series, pareto_info: dict = None) -> str:
    """Generate domain-specific engineering commentary mapped to test numbers (Tno)."""
    lot_no = row.get("LOT_No", "N/A")
    yield_val = row.get("yield_rate", 0.0)

    top_tno = pareto_info.get("top_tno", "N/A") if pareto_info else "N/A"
    top_count = pareto_info.get("top_count", 0) if pareto_info else 0

    # Professional diagnostic rule mapping based on test number (Tno) ranges
    if top_tno != "N/A":
        issue_desc = f"最頻不具合項目: {top_tno} ({top_count} pcs)"
        tno_digits = "".join(filter(str.isdigit, top_tno))
        tno_num = int(tno_digits) if tno_digits else 0

        if 100 <= tno_num < 200:
            advice = (
                "アナログ検査項目での規格外れが多発。"
                "コンパレータ基準電圧の変動や入力アンプのオフセット特性を確認してください。"
            )
        elif 200 <= tno_num < 300:
            advice = (
                "コンタクトオープン不良が主因。"
                "プローブカードの接触針圧低下、ピッチズレ、あるいはパッド表面の酸化皮膜を点検してください。"
            )
        elif 300 <= tno_num < 500:
            advice = (
                "リーク電流テスト不良が多発。"
                "ESD保護素子の破壊、あるいはゲート酸化膜の局所欠陥リークを調査してください。"
            )
        elif 500 <= tno_num < 600:
            advice = (
                "動作/静止電流（Icc）不良が多発。"
                "電源ラインの微小ショートや内部クロックのリークパスを検証してください。"
            )
        elif 600 <= tno_num < 800:
            advice = (
                "入出力電流ドライブ能力不足。"
                "出力トランジスタのサイズバリエーションやドライバ段の電圧降下を調査してください。"
            )
        elif 800 <= tno_num < 950:
            advice = (
                "ファンクションテスト（動作マージン）不良。"
                "VMIN/VMAX動作条件におけるタイミングクリティカルパスと電源デカップリングを確認してください。"
            )
        else:
            advice = (
                f"該当テスト項目 ({top_tno}) の判定閾値ばらつき抑制および"
                "テストシーケンスの再点検を推奨します。"
            )
    else:
        issue_desc = "主要不具合項目: 検出なし"
        advice = "現行の製造条件を維持し、トレンド監視を継続してください。"

    if yield_val >= 99.0:
        eval_text = "目標に対し非常に良好な歩留まりを維持しています。"
    elif yield_val >= 95.0:
        eval_text = (
            "概ね安定していますが、特定Tnoによる軽微な歩留まりロスが発生しています。"
        )
    else:
        eval_text = (
            "目標歩留まりを下回っており（Yield Loss過大）、早急なプロセス改善が必要です。"
        )

    return (
        f"• 対象LOT: {lot_no} (実績歩留まり: {yield_val:.2f}%)\n"
        f"• 状況評価: {eval_text}\n"
        f"• {issue_desc}\n"
        f"• 推奨対策: {advice}"
    )


# ---------------------------------------------------------
# 4. Automated PowerPoint Report Generator
# ---------------------------------------------------------
def create_powerpoint_report(
    df: pd.DataFrame,
    analysis: dict,
    trend_img: Path,
    pareto_results: list,
    output_pptx: Path,
):
    """Generate customized widescreen PowerPoint presentation report."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title slide
    slide = prs.slides.add_slide(blank_layout)
    tb = slide.shapes.add_textbox(
        Inches(1.0), Inches(2.5), Inches(11.333), Inches(2.0)
    )
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "半導体システムLSI 歩留まり分析レポート"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    p2 = tf.add_paragraph()
    p2.text = (
        f"生成日時: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} | "
        "製造技術部・品質保証部（サンプル）"
    )
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(100, 116, 139)

    # Slide 2: Overall yield trend slide
    slide = prs.slides.add_slide(blank_layout)
    tb = slide.shapes.add_textbox(
        Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8)
    )
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "1. 歩留まりトレンド"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    if trend_img.exists():
        slide.shapes.add_picture(
            str(trend_img), Inches(0.8), Inches(1.6), width=Inches(6.5)
        )

    tb_desc = slide.shapes.add_textbox(
        Inches(7.6), Inches(1.6), Inches(4.9), Inches(4.5)
    )
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    p_d1 = tf_desc.paragraphs[0]
    p_d1.text = "【トレンド分析結果】"
    p_d1.font.bold = True
    p_d1.font.size = Pt(20)

    p_d2 = tf_desc.add_paragraph()
    p_d2.text = (
        f"• {analysis['overall_status']}\n"
        f"• ワーストLOT: 「{analysis['worst_lot_no']}」 "
        f"({analysis['worst_yield']:.2f}%)\n\n"
        f"• {analysis['trend_comparison']}"
    )
    p_d2.font.size = Pt(14)
    p_d2.space_before = Pt(10)

    # Slide 3+: Individual lot Pareto analysis slides
    for item in pareto_results:
        lot_no = item["lot_no"]
        p_img = item["image_path"]

        matched_rows = df[df["LOT_No"].astype(str).str.strip() == str(lot_no)]
        row_data = (
            matched_rows.iloc[0]
            if not matched_rows.empty
            else pd.Series({"LOT_No": lot_no, "yield_rate": 0.0})
        )

        lot_analysis_text = get_lot_specific_analysis(
            row_data, pareto_info=item
        )

        slide = prs.slides.add_slide(blank_layout)

        # Slide header title
        tb = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.3), Inches(11.7), Inches(0.8)
        )
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"2. 不良要因 パレート分析 LOT_NO.{lot_no}"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)

        # Top section: Pareto chart image
        if p_img.exists():
            slide.shapes.add_picture(
                str(p_img), Inches(0.8), Inches(1.1), width=Inches(11.7)
            )

        # Bottom section: Engineering analysis text box
        tb_desc2 = slide.shapes.add_textbox(
            Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.8)
        )
        tf_desc2 = tb_desc2.text_frame
        tf_desc2.word_wrap = True

        p_c1 = tf_desc2.paragraphs[0]
        p_c1.text = "【プロセス解析・改善アクション】"
        p_c1.font.bold = True
        p_c1.font.size = Pt(18)

        p_c2 = tf_desc2.add_paragraph()
        p_c2.text = lot_analysis_text
        p_c2.font.size = Pt(13)
        p_c2.space_before = Pt(5)

    prs.save(str(output_pptx))
    logger.info(f"Successfully exported PowerPoint report: {output_pptx}")


# ---------------------------------------------------------
# Temporary Image Cleanup
# ---------------------------------------------------------
def cleanup_temp_images(image_paths: list):
    """Clean up and remove temporary PNG charts generated for PowerPoint insertion."""
    for img_path in image_paths:
        try:
            p = Path(img_path)
            if p.exists():
                p.unlink()
                logger.info(f"Deleted temporary image: {p.name}")
        except Exception as e:
            logger.warning(f"Failed to delete temporary image ({img_path}): {e}")


# ---------------------------------------------------------
# 5. Batch Pipeline Logic
# ---------------------------------------------------------
def process_batch():
    """Main execution pipeline for batch processing input Excel files."""
    target_files = glob.glob(str(INPUT_DIR / "*.xlsx")) + glob.glob(
        str(INPUT_DIR / "*.xls")
    )

    if not target_files:
        print(f"[警告] '{INPUT_DIR}' フォルダにExcelファイルが見つかりません。")
        return

    for file_path in target_files:
        path = Path(file_path)
        if path.name.startswith("~$"):
            continue

        print(f"-> 処理中: {path.name}")
        try:
            df = load_and_validate_yield_excel(path)
            if df.empty:
                print(
                    f"   [警告] 有効なデータが取得できませんでした: {path.name}\n"
                )
                continue

            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            trend_img = OUTPUT_DIR / f"trend_{path.stem}_{timestamp}.png"

            # 1. Generate chart images
            generate_trend_chart(df, trend_img)
            pareto_results = generate_lot_pareto_charts(
                path, OUTPUT_DIR, path.stem, timestamp
            )

            # 2. Build and save PowerPoint report
            analysis = get_llm_analysis_summary(df)
            output_pptx = (
                OUTPUT_DIR
                / f"Yield_Analysis_Report_{path.stem}_{timestamp}.pptx"
            )
            create_powerpoint_report(
                df, analysis, trend_img, pareto_results, output_pptx
            )

            # 3. Clean up temporary chart images
            temp_images = [trend_img] + [
                item["image_path"] for item in pareto_results
            ]
            cleanup_temp_images(temp_images)

            print(
                f"   [成功] レポート出力完了（一時画像削除済）: outputs/{output_pptx.name}\n"
            )
        except Exception as e:
            print(f"   [エラー] {path.name} の処理に失敗しました: {e}\n")


if __name__ == "__main__":
    print("=== Start Semiconductor LSI Yield Analysis Report Generation ===")
    process_batch()
    print("=== Finished ===")